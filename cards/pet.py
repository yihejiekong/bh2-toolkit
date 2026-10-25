# -*- coding: utf-8 -*-
"""使魔（宠物）图鉴卡生成器 —— 沿用武器卡模板，改动三处：

  1. 属性块：攻击力（满级伤害）/ 暴击率（攻速不出图）
  2. 属性块下一行：该使魔持有的词条，格式「图标 词条名」，字号与属性块一致，颜色用浅灰
  3. 技能名下方、技能描述上方：灰色小一号字写「初始CD X秒　CD Y秒」；技能名右侧带类型标签
     （奥义 / 奥义·觉醒 / 自律 / 固有）

数据来源：xlsx 的「宠物」sheet（239 行 × 71 列，skilldata 已合并）。
用法：python pet.py 971 [更多游戏内ID...]
"""

import math
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import cardgen as C
import paths as P

SHEET = '宠物'
LABEL = '使魔'
TRAIT_DIR = os.path.join(os.path.join(P.MATERIALS, '词条图标'), 'flat')


def load_pet():
    wb = openpyxl.load_workbook(C.XLSX, read_only=True)
    ws = wb[SHEET]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(hdr)}
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        dn = str(row[idx['游戏内ID']])
        rarity = C.num(row[idx['稀有度']])
        cur = rows.get(dn)
        if cur is None or rarity > C.num(cur[idx['稀有度']]):
            rows[dn] = row
    return idx, rows


def load_traits():
    wb = openpyxl.load_workbook(C.XLSX, read_only=True)
    ws = wb['词条']
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        out[int(row[0])] = (row[1] or '', C.clean(row[8] or ''))
    return out


IDX, ROWS = load_pet()
TRAITS = load_traits()


def pct(v):
    v = C.num(v) * 100
    return C.trim(v) + '%'


def floor1(v):
    return math.floor(v + 1e-6)


def skill_list(row, maxlv):
    """返回 [(名称, 描述, meta, 类型标签)]。

    初始CD 解包表里没有字段（petskilldata 只有 coolDown/coolDownAdd/cooldownSkillUpAdd/
    cooldownIsReduced），实测游戏内固定为 CD 的 90% 向下取整（壳之幼子 40→36、
    天之幼子 40→36、芝士 45→40.5→40），故按此推算。
    """
    name2meta = {}
    for key, cd_key, kind in (('奥义名称', '奥义CD', '奥义'),
                              ('奥义·觉醒名称', '奥义·觉醒CD', '奥义·觉醒'),
                              ('自律名称', '自律CD', '自律'),
                              ('固有名称', '固有CD', '固有')):
        nm = row[IDX[key]]
        if not nm:
            continue
        cd = C.num(row[IDX[cd_key]])
        if kind in ('奥义', '奥义·觉醒'):
            parts = ['初始CD %s秒' % C.trim(floor1(cd * 0.9)), 'CD %s秒' % C.trim(cd)]
            name2meta[str(nm)] = ('　'.join(parts), kind)
        else:
            meta = ('CD %s秒' % C.trim(cd)) if cd else None
            name2meta[str(nm)] = (meta, kind)

    out = []
    for t_, b_ in C.split_skills(row[IDX['技能（解包原文）']] or ''):
        meta, kind = name2meta.get(t_, (None, None))
        out.append((t_, C.expand_params(b_, maxlv), meta, kind))
    return out


def build(dn):
    row = ROWS[str(dn)]
    g = lambda k: row[IDX[k]]
    maxlv = int(C.num(g('最大等级'), 1))

    stats = [('攻击力', C.trim(C.num(g('满级伤害')))),
             ('暴击率', pct(g('暴击率')))]

    traits = []
    for tid in str(g('词条ID') or '').split(';'):
        tid = tid.strip()
        if not tid or tid == '0':
            continue
        name = TRAITS.get(int(tid), ('', ''))[0]
        f = '%s_%s.png' % (tid, name)
        src = os.path.join(TRAIT_DIR, f)
        if not os.path.exists(src):
            continue
        dest = os.path.join(C.ASSETS, f)
        if not os.path.exists(dest):
            shutil.copyfile(src, dest)
        traits.append({'name': name, 'file': f})

    icon_file = g('图标文件')
    if not icon_file:
        icon_file = C.equip_icon(g('图标ID'))
    src = os.path.join(C.MDIR, '图标', icon_file or '')
    dest = os.path.join(C.ASSETS, icon_file or '')
    if icon_file and os.path.exists(src) and not os.path.exists(dest):
        shutil.copyfile(src, dest)

    return {
        'name': g('名称'),
        'rarity': int(C.num(g('稀有度'))),
        'type': LABEL,
        'series': g('系列') or '',
        'maxlv': maxlv,
        'stats': stats,
        'traits': traits,
        'skills': skill_list(row, maxlv),
        'note': C.clean(g('描述（解包原文）') or ''),
        'icon': icon_file,
        'figure': '',
        'faces': [],
        'mats': [],
        'bio': None,
    }


def main(dns):
    for dn in dns:
        card = build(dn)
        path = os.path.join(C.OUT, 'card_%s.html' % dn)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(C.render(card))
        print('wrote', path, os.path.getsize(path),
              '| 属性:', card['stats'],
              '| 词条:', [t['name'] for t in card['traits']],
              '| 技能:', [(t, k, m) for t, _, m, k in card['skills']])


if __name__ == '__main__':
    main(sys.argv[1:] or ['971'])


