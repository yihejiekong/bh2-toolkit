# -*- coding: utf-8 -*-
"""徽章（被动技能装备）图鉴卡生成器 —— 与服装卡同一模板，属性块只放「负重 / 亲密度上限」。

数据来源：xlsx 的「使魔-被动技能」sheet（对应解包表 passiveskilldatav3，1867 行）。
差异点（其余完全沿用 cardgen.py 的渲染）：
  · 属性块：负重 + 亲密度上限（徽章没有生命值，也没有攻击/载弹/攻速/暴击）
  · LV 后面的类型文字固定写「徽章」（表内「类别」列写的是 使魔/被动技能）
  · 觉醒块 / 系列胶囊 / 觉醒素材 / 技能数值展开 完全一致
用法：python badge.py 2853 [更多游戏内ID...]
"""

import glob
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import cardgen as C
import paths as P

SHEET = '使魔-被动技能'
LABEL = '徽章'


def load_badge():
    wb = openpyxl.load_workbook(C.XLSX, read_only=True)
    ws = wb[SHEET]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(hdr)}
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        dn = str(row[idx['游戏内ID']])
        rarity = C.num(row[idx['稀有度']])
        cur = rows.get(dn)
        if cur is None or rarity > C.num(cur[idx['稀有度']]):
            rows[dn] = row
    return idx, rows


IDX, ROWS = load_badge()


def build(dn):
    row = ROWS[str(dn)]
    g = lambda k: row[IDX[k]]
    maxlv = int(C.num(g('最大等级'), 1))

    stats = [('负重', C.trim(C.num(g('负重'))))]
    if C.num(g('亲密度上限')) > 0:
        stats.append(('亲密度上限', C.trim(C.num(g('亲密度上限')))))

    icon_file = g('图标文件')
    fig_file = g('觉醒立绘文件')
    pid = int(C.num(g('立绘ID')))
    if not icon_file:
        icon_file = C.equip_icon(g('图标ID'))
    fig_small = False
    if not fig_file and pid:
        # 立绘ID 是 PosterID，必须换成 FigureID 再找图
        fid = C.POSTER_FIG.get(pid)
        if fid:
            cand = glob.glob(os.path.join(C.MDIR, '觉醒立绘', '%d_*.png' % fid))
            if cand:
                fig_file = os.path.basename(sorted(cand)[0])
            else:
                small = os.path.join(C.MDIR, '看板小图', '%d.png' % fid)
                if os.path.exists(small):
                    fig_file = 'small_%d.png' % fid
                    fig_small = True
                    shutil.copyfile(small, os.path.join(C.ASSETS, fig_file))
    for src, dest in ((os.path.join(C.MDIR, '图标', icon_file or ''), os.path.join(C.ASSETS, icon_file or '')),
                      (os.path.join(C.MDIR, '觉醒立绘', fig_file or ''), os.path.join(C.ASSETS, fig_file or ''))):
        if src and os.path.exists(src) and not os.path.exists(dest):
            shutil.copyfile(src, dest)
    fig_id = os.path.basename(fig_file).split('_')[0] if fig_file else ''
    faces = C.face_files(fig_id) if fig_file else []

    mats = []
    for mid in C.PERSONA.get(int(C.num(g('内部ID'))), []):
        name, icon_id = C.MAT.get(mid, ('', ''))
        f = C.mat_icon(icon_id)
        if f:
            mats.append({'name': name, 'file': f, 'count': 1})
    merged = []
    for m in mats:
        for n in merged:
            if n['name'] == m['name']:
                n['count'] += 1
                break
        else:
            merged.append(m)
    mats = merged

    return {
        'name': g('名称'),
        'rarity': int(C.num(g('稀有度'))),
        'type': LABEL,
        'series': g('系列') or '',
        'maxlv': maxlv,
        'stats': stats,
        'skills': [(t_, C.expand_params(b_, maxlv))
                   for t_, b_ in C.split_skills(g('技能（解包原文）'))],
        'note': C.clean(g('描述（解包原文）') or ''),
        'icon': icon_file,
        'figure': fig_file,
        'fig_small': fig_small,
        'faces': faces,
        'mats': mats,
        'bio': C.poster_bio(pid),
    }


def main(dns):
    for dn in dns:
        card = build(dn)
        path = os.path.join(C.OUT, 'card_%s.html' % dn)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(C.render(card))
        print('wrote', path, os.path.getsize(path),
              '| 属性:', [k for k, _ in card['stats']],
              '| 觉醒立绘:', card['figure'],
              '| 表情:', len(card['faces']),
              '| 材料:', ['%s×%d' % (m['name'], m['count']) for m in card['mats']])


if __name__ == '__main__':
    main(sys.argv[1:] or ['2853'])
